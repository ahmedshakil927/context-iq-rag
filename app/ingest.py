import io
from docx import Document
from pypdf import PdfReader
from .config import CHUNK_SIZE, CHUNK_OVERLAP


def extract_pages(filename: str, data: bytes) -> list[tuple[int, str]]:
    """Return (page_number, text) pairs. Plain text files count as page 1."""
    if filename.lower().endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        return [(i + 1, p.extract_text() or "") for i, p in enumerate(reader.pages)]
    if filename.lower().endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for t in doc.tables:
            parts += [" | ".join(c.text for c in row.cells) for row in t.rows]
        return [(1, "\n".join(parts))]
    return [(1, data.decode("utf-8", errors="ignore"))]


def chunk_text(text: str, size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> list[str]:
    text = " ".join(text.split())
    chunks, start = [], 0
    while start < len(text):
        chunks.append(text[start:start + size])
        start += size - overlap
    return [c for c in chunks if c.strip()]
