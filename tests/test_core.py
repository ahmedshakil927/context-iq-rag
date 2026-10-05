"""Fast tests: no Ollama needed (embeddings and generation are faked)."""
import io
import math

import numpy as np
import pytest
from docx import Document
from fastapi.testclient import TestClient

from app import ingest, llm, rag, store as store_mod
from app.store import VectorStore

DIM = 64


def fake_embed(texts):
    """Bag-of-words hashing embedding: similar words -> similar vectors."""
    out = []
    for t in texts:
        v = np.zeros(DIM)
        for w in store_mod.tokenize(t):
            v[hash(w) % DIM] += 1
        out.append((v / (np.linalg.norm(v) + 1e-9)).tolist())
    return out


@pytest.fixture(autouse=True)
def fakes(monkeypatch, tmp_path):
    monkeypatch.setattr(llm, "embed", fake_embed)
    monkeypatch.setattr(llm, "generate", lambda prompt, model=None: "FAKE ANSWER")
    monkeypatch.setattr(llm, "generate_stream", lambda prompt, model=None: iter(["FAKE ", "ANSWER"]))
    monkeypatch.setattr(store_mod, "STORE_PATH", str(tmp_path / "store.json"))


# ---------- ingest ----------
def test_chunking_respects_size_and_overlap():
    text = "word " * 1000
    chunks = ingest.chunk_text(text, size=200, overlap=50)
    assert all(len(c) <= 200 for c in chunks)
    assert len(chunks) > 20
    assert chunks[0][-30:].strip() in chunks[1]          # overlap carries text across chunks


def test_chunking_empty_text():
    assert ingest.chunk_text("   \n  ") == []


def test_docx_text_and_tables_are_extracted():
    d = Document()
    d.add_paragraph("Hello paragraph")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Lead", "Priya"
    buf = io.BytesIO()
    d.save(buf)
    (page, text), = ingest.extract_pages("x.docx", buf.getvalue())
    assert "Hello paragraph" in text and "Priya" in text


# ---------- store ----------
def make_store(rows):
    s = VectorStore(load=False)
    for doc, text in rows:
        s.add(doc, 1, text, fake_embed([text])[0])
    return s


def test_vector_search_finds_most_similar():
    s = make_store([("a", "cats purr and sleep"), ("b", "stock markets fell sharply today")])
    q = "why do cats purr"
    hits, best, _ = s.search(fake_embed([q])[0], q, 1)
    assert hits[0][0]["doc"] == "a" and best > 0.3


def test_hybrid_ranks_exact_keyword_match_first():
    rows = [("a", f"generic filler sentence number {i} about nothing") for i in range(20)]
    rows.append(("needle", "the Zorblax engine runs at 4200 rpm"))
    s = make_store(rows)
    q = "Zorblax rpm"
    hits, _, kw = s.search(fake_embed([q])[0], q, 3, "hybrid")
    assert "needle" in [h[0]["doc"] for h in hits]
    assert kw[1] == 2                                     # both content words matched in one chunk


def test_remove_and_documents_listing():
    s = make_store([("a", "one"), ("a", "two"), ("b", "three")])
    assert [d["name"] for d in s.documents()] == ["a", "b"]
    assert s.remove("a") == 2
    assert [d["name"] for d in s.documents()] == ["b"]


# ---------- relevance gate ----------
def test_keyword_gate_accepts_rare_term_when_vector_score_is_low(monkeypatch):
    s = make_store([("doc", "The Zorblax engine runs at 4200 rpm and was invented by Dr. Quill")])
    monkeypatch.setattr(llm, "embed", lambda texts: [[1.0] + [0.0] * (DIM - 1)] * len(texts))  # orthogonal -> low vector score
    q = "Who invented the Zorblax engine?"
    assert rag.retrieve(s, q, kw_gate=False)[0] is None   # old behaviour: refused
    assert rag.retrieve(s, q, kw_gate=True)[0] is not None  # keyword check lets it through


def test_gate_still_refuses_unrelated_question(monkeypatch):
    s = make_store([("doc", "The Zorblax engine runs at 4200 rpm")])
    monkeypatch.setattr(llm, "embed", lambda texts: [[1.0] + [0.0] * (DIM - 1)] * len(texts))
    assert rag.retrieve(s, "What is the capital of France?", kw_gate=True)[0] is None


# ---------- API ----------
@pytest.fixture
def client(monkeypatch):
    from app import main
    monkeypatch.setattr(main, "store", VectorStore(load=False))
    return TestClient(main.app)


def upload(client, name, content):
    return client.post("/upload", files={"file": (name, content)})


def test_upload_rejects_unsupported_and_empty(client):
    assert upload(client, "a.exe", b"x").status_code == 415
    assert upload(client, "empty.txt", b"").status_code == 422


def test_upload_ask_delete_flow(client):
    assert upload(client, "n.txt", b"Project Nimbus launched in March 2031.").json()["chunks"] == 1
    assert client.get("/documents").json()[0]["name"] == "n.txt"
    r = client.post("/ask", json={"question": "When did Project Nimbus launch?"}).json()
    assert r["answer"] == "FAKE ANSWER" and r["sources"][0]["doc"] == "n.txt"
    assert client.delete("/documents/n.txt").status_code == 200
    assert client.delete("/documents/n.txt").status_code == 404
    assert client.get("/documents").json() == []


def test_reupload_replaces_instead_of_duplicating(client):
    upload(client, "n.txt", b"first version of the text")
    upload(client, "n.txt", b"second version of the text")
    docs = client.get("/documents").json()
    assert len(docs) == 1 and docs[0]["chunks"] == 1


def test_ask_with_nothing_indexed_refuses(client):
    r = client.post("/ask", json={"question": "anything?"}).json()
    assert r["answer"] == rag.REFUSAL and r["sources"] == []


def test_stream_endpoint_emits_sources_tokens_done(client):
    upload(client, "n.txt", b"Project Nimbus launched in March 2031.")
    lines = client.post("/ask/stream", json={"question": "When did Project Nimbus launch?"}).text.strip().split("\n")
    types = [__import__("json").loads(l)["type"] for l in lines]
    assert types[0] == "sources" and types[-1] == "done" and "token" in types
